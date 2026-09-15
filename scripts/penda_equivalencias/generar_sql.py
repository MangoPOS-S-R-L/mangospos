#!/usr/bin/env python3
"""Genera el SQL que carga las equivalencias de La Penda (1 unidad = N g / mL)
desde la hoja «Equivalencias» de docs/penda_recetas/REVISION_RECETARIO_PENDA.xlsx.

No toca la base: escribe dos archivos .sql para correr en Studio.

CANDADOS
  En la hoja:
    - Solo toma filas con «Confirmado por cocina» = sí, un número > 0 y unidad
      g o mL. Lo demás se lista como descartado, con el motivo.
  En el SQL:
    - Toca SOLO el negocio de La Penda y SOLO los insumos por id: un id de
      otro negocio no hace nada, así que ningún otro negocio —tenga insumos o
      no— queda expuesto.
    - Nunca pisa una equivalencia ya cargada (`conversion_unit is null`).
    - Salta un insumo cuya unidad base ya es una medida (g, lb, mL…): ahí la
      equivalencia no aplica.
    - Aborta sin cargar nada si falta la migración 20260915_0001.
    - Tope de bloqueo, igual que la migración.
    - Deja un respaldo con lo cargado, sin acceso desde la API. El ROLLBACK
      solo borra lo que sigue igual a lo cargado: no deshace una equivalencia
      que alguien corrigió después desde la app.

Uso:
    python3 scripts/penda_equivalencias/generar_sql.py [hoja.xlsx] [carpeta_salida]
Requiere openpyxl.
"""

import datetime
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

PENDA = '35c5076a-bd85-4a1b-8d1c-ce719c4f9ae6'
SHEET = 'Equivalencias'
BACKUP = 'public.backup_penda_equivalencias_20260915'
UUID_RE = re.compile(r'^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$')
# Unidad de la hoja → código del catálogo de la app (unit_catalog.dart).
UNITS = {'g': 'g', 'gr': 'g', 'gramos': 'g', 'ml': 'ml', 'mililitros': 'ml'}
YES = {'si', 'sí', 's', 'x', 'yes', 'ok'}

REPO = Path(__file__).resolve().parents[2]
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / 'docs/penda_recetas/REVISION_RECETARIO_PENDA.xlsx'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else REPO / 'docs/penda_recetas'


def col(headers, prefix):
    for i, h in enumerate(headers):
        if h and str(h).strip().lower().startswith(prefix.lower()):
            return i
    raise SystemExit(f'La hoja «{SHEET}» no tiene la columna «{prefix}…»')


def number(value):
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(',', '.'))
    except ValueError:
        return None


def sql_text(s):
    return "'" + str(s).replace("'", "''") + "'"


def main():
    ws = load_workbook(XLSX, data_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    c_name, c_val = col(headers, 'Insumo'), col(headers, '1 [unidad hoy]')
    c_unit, c_ok = col(headers, 'Unidad (g'), col(headers, 'Confirmado')
    c_id = col(headers, 'ID')

    load, skipped = [], []
    for r in rows[1:]:
        name = r[c_name]
        if not name:
            continue
        confirmed = str(r[c_ok] or '').strip().lower() in YES
        value, unit = number(r[c_val]), UNITS.get(str(r[c_unit] or '').strip().lower())
        item_id = str(r[c_id] or '').strip().lower()
        if not confirmed:
            reason = 'sin confirmar'
        elif value is None or value <= 0:
            reason = 'confirmado pero sin número válido'
        elif unit is None:
            reason = f'unidad «{r[c_unit]}» no es g ni mL'
        elif not UUID_RE.match(item_id):
            reason = 'ID vacío o inválido'
        else:
            load.append((item_id, unit, value, name))
            continue
        skipped.append((name, reason))

    ids = [x[0] for x in load]
    if len(ids) != len(set(ids)):
        raise SystemExit('Hay un mismo insumo repetido en la hoja: corregir antes de generar.')

    stamp = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    values = ',\n'.join(
        f"    ({sql_text(i)}::uuid, {sql_text(u)}, {v!r}::numeric, {sql_text(n)})"
        for i, u, v, n in load
    )
    measures = "('g','gr','kg','lb','lbs','oz','ml','l','lt','fl oz','gal','qt','cup','tbsp','tsp','mg','cl')"

    carga = f"""-- =============================================================================
-- CARGAR equivalencias de La Penda — generado {stamp}
-- Desde: {XLSX.name} › {SHEET}. {len(load)} insumos confirmados por cocina.
--
-- Solo toca el negocio {PENDA} y solo los insumos por id.
-- No pisa equivalencias ya cargadas ni toca insumos cuya base ya es una medida.
-- Deshacer: CARGAR_EQUIVALENCIAS_PENDA_ROLLBACK.sql
-- =============================================================================

begin;

set local lock_timeout = '5s';
set local statement_timeout = '60s';

do $$
begin
  if not exists (
    select 1 from information_schema.columns
     where table_schema = 'public'
       and table_name   = 'inventory_items'
       and column_name  = 'conversion_unit'
  ) then
    raise exception 'Falta aplicar la migración 20260915_0001_inventory_item_unit_conversion. No se cargó nada.';
  end if;
end $$;

-- Respaldo de lo que se carga. Sin acceso desde la API (anon/authenticated).
create table if not exists {BACKUP} (
  id             uuid        not null,
  business_id    uuid        not null,
  nombre         text,
  cargado_unit   text        not null,
  cargado_factor numeric     not null,
  cargado_en     timestamptz not null default now()
);
revoke all on table {BACKUP} from anon, authenticated;

create temporary table tmp_equivalencias (
  id uuid, conversion_unit text, conversion_factor numeric, nombre text
) on commit drop;

insert into tmp_equivalencias (id, conversion_unit, conversion_factor, nombre) values
{values if values else "    (null, null, null, null)"};

create temporary table tmp_objetivo on commit drop as
select ii.id, t.conversion_unit, t.conversion_factor, ii.name
  from tmp_equivalencias t
  join public.inventory_items ii on ii.id = t.id
 where ii.business_id = '{PENDA}'::uuid
   and ii.conversion_unit is null
   and lower(coalesce(ii.unit, 'unidad')) not in {measures};

insert into {BACKUP} (id, business_id, nombre, cargado_unit, cargado_factor)
select o.id, '{PENDA}'::uuid, o.name, o.conversion_unit, o.conversion_factor
  from tmp_objetivo o;

update public.inventory_items ii
   set conversion_unit   = o.conversion_unit,
       conversion_factor = o.conversion_factor
  from tmp_objetivo o
 where ii.id = o.id
   and ii.business_id = '{PENDA}'::uuid
   and ii.conversion_unit is null;

-- Qué se cargó y qué NO (con el motivo). Revisar ANTES del commit.
select coalesce(ii.name, t.nombre) as insumo,
       t.conversion_unit, t.conversion_factor,
       case
         when ii.id is null                          then 'NO: el id no existe'
         when ii.business_id <> '{PENDA}'::uuid       then 'NO: el insumo es de OTRO negocio'
         when o.id is not null                       then 'cargado'
         when ii.conversion_unit is not null         then 'NO: ya tenía equivalencia'
         else 'NO: su unidad base ya es una medida'
       end as resultado
  from tmp_equivalencias t
  left join public.inventory_items ii on ii.id = t.id
  left join tmp_objetivo o on o.id = t.id
 where t.id is not null
 order by resultado, insumo;

commit;
"""

    rollback = f"""-- =============================================================================
-- ROLLBACK de CARGAR_EQUIVALENCIAS_PENDA.sql — generado {stamp}
--
-- Borra SOLO las equivalencias que siguen igual a como se cargaron: si alguien
-- corrigió una después desde la app, esa se respeta.
-- =============================================================================

begin;

set local lock_timeout = '5s';
set local statement_timeout = '60s';

update public.inventory_items ii
   set conversion_unit = null,
       conversion_factor = null
  from {BACKUP} b
 where ii.id = b.id
   and ii.business_id = b.business_id
   and ii.business_id = '{PENDA}'::uuid
   and ii.conversion_unit = b.cargado_unit
   and ii.conversion_factor = b.cargado_factor;

commit;
"""

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'CARGAR_EQUIVALENCIAS_PENDA.sql').write_text(carga, encoding='utf-8')
    (OUT / 'CARGAR_EQUIVALENCIAS_PENDA_ROLLBACK.sql').write_text(rollback, encoding='utf-8')

    print(f'Cargables: {len(load)}')
    for i, u, v, n in load:
        print(f'  1 = {v:g} {u:<2}  {n}')
    print(f'Descartadas: {len(skipped)}')
    for n, why in skipped:
        print(f'  {why:<34} {n}')
    print(f'SQL en {OUT}')


if __name__ == '__main__':
    main()
