#!/usr/bin/env python3
"""
Estancia Nueva Sports Club (RNC 133231468) — división del inventario del sistema
viejo en CAFETERÍA y TIENDA.

Fuente: ~/Desktop/ARTICULOS DE INVENTARIO.csv ("Listar Existencia de Artículos",
1,015 renglones: código, descripción, último costo, existencia, costo total).
El archivo NO trae precio de venta ni categorías.

Salidas (en esta misma carpeta):
  CAFETERIA_Estancia_Nueva.xlsx  hoja 1 "Productos" en el formato de
                                 Productos → Importar catálogo (falta llenar
                                 PRECIO), + Insumos, Revisar y Resumen.
  TIENDA_Estancia_Nueva.xlsx     solo dividido (pádel, fútbol, ropa, servicios).
  POR_DECIDIR_Estancia_Nueva.xlsx  lo que no es claramente de ninguno de los dos.

Regenerar:  <python con openpyxl> build_estancia_nueva.py
Editar ESTE archivo, no los .xlsx.
"""
import collections
import csv
import os
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = os.path.expanduser('~/Desktop/ARTICULOS DE INVENTARIO.csv')
BID = '85924083-2e8e-4e64-8192-808ee24674ed'  # negocio de la CAFETERÍA
OUT = os.path.dirname(os.path.abspath(__file__))


# --------------------------------------------------------------------------
# Lectura
# --------------------------------------------------------------------------
def num(s):
    """' 1,285.00 ' → 1285.0 · ' -   ' → 0 · '(3.00)' → -3.0"""
    s = (s or '').strip()
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()').replace(',', '').strip()
    if s in ('', '-'):
        return 0.0
    return -float(s) if neg else float(s)


def load():
    with open(SRC, encoding='utf-8-sig', newline='') as f:
        table = list(csv.reader(f))
    hdr = next(i for i, x in enumerate(table) if x and x[0] == 'Referencia')
    rows = []
    for i, x in enumerate(table[hdr + 1:], start=hdr + 2):
        if not any(v.strip() for v in x):
            continue
        rows.append(dict(line=i, code=x[1].strip(), name=x[2].strip(),
                         cost=num(x[3]), stock=num(x[4]), unit=x[5].strip(),
                         total=num(x[6])))
    return rows


# --------------------------------------------------------------------------
# Códigos de barra: solo GTIN que validan (EAN-8/13, UPC-A) o UPC-E.
# Los internos (4 dígitos, 110000100xxx) y los rotos quedan solo como SKU.
# --------------------------------------------------------------------------
def gtin_ok(c):
    if not c.isdigit() or len(c) not in (8, 12, 13, 14):
        return False
    d = [int(x) for x in c]
    s = sum(v * (3 if i % 2 == 0 else 1) for i, v in enumerate(reversed(d[:-1])))
    return (10 - s % 10) % 10 == d[-1]


def upce_ok(c):
    if len(c) != 8 or not c.isdigit() or c[0] not in '01':
        return False
    d, x = c[1:7], c[6]
    if x in '012':
        body = d[0:2] + x + '0000' + d[2:5]
    elif x == '3':
        body = d[0:3] + '00000' + d[3:5]
    elif x == '4':
        body = d[0:4] + '00000' + d[4]
    else:
        body = d[0:5] + '0000' + x
    return gtin_ok(c[0] + body + c[7])


def barcode_for(code):
    if code.startswith('1100001'):
        return ''
    return code if (gtin_ok(code) or upce_ok(code)) else ''


# --------------------------------------------------------------------------
# Clasificación. (patrón, destino, grupo/categoría, tipo o motivo).
# Primer match gana. Destino: C = cafetería, T = tienda, X = por decidir.
# --------------------------------------------------------------------------
def R(p):
    return re.compile(p)


RULES = [
  # ---- Por decidir -------------------------------------------------------
  (R(r'^PRUEBA CODIGO'), 'X', 'Basura', 'Prueba del sistema viejo'),
  (R(r'^14$'), 'X', 'Basura', 'Nombre roto ("14")'),
  (R(r'^NOVA$'), 'X', 'Desconocido', 'No se sabe qué es (costo 1.00)'),
  (R(r'^HUACAL$|^BOTELLAS VACIA'), 'X', 'Envases', 'Envase/depósito de cerveza: no se vende'),
  (R(r'ALIVIOL|ONTOL|ONTO GEL|CURA CLEAN|WINAZORD'), 'X', 'Botiquín', 'Medicamento: ¿cafetería o tienda?'),
  (R(r'CUMPLEA'), 'X', 'Eventos', 'Paquete de cumpleaños: ¿cafetería o club?'),
  (R(r'NEVERITA DESECHABLE'), 'X', 'Otros', '¿Se vende en la barra o en la tienda?'),
  (R(r'^FUNPACK$'), 'X', 'Desconocido', 'No se sabe qué es (¿empaque? ¿dulce?)'),
  # ---- Tienda: servicios del club ----------------------------------------
  (R(r'CLASES? DE +PADEL|ALQUILER DE PALAS|LIGA ABIERTA|PATROCINIO|VALLAS? LATERALES|SERVICIOS DE PUBLICIDAD|SERVICIO DE TRANSPORTE'), 'T', 'Servicios del club', ''),
  # ---- Tienda: pádel -----------------------------------------------------
  (R(r'\bPALA\b|PALETERO|GRIPS?\b|OVERGRIP|PELOTAS?\b|BOLAS DE PADEL|BABOLAT|BULLPADEL|ARROW HIT|PADEL RUSH|PRO NOX|AT10|MUÑEQUERA|PROTECTOR SIUX|CODERA|VIBOR-?A|SACK PACK|MOCHILA SOFTEE'), 'T', 'Pádel', ''),
  # ---- Tienda: fútbol ----------------------------------------------------
  (R(r'BALON|CAMISETA|CONJU?N?TO|ESPINILLERA|SPINILLERA|GUANTE|ZAPATILLA|BOTAS|GORRA|MOCHILA|PERFUME BARCA|INFLADOR|KIT CONOS|PIZARRA|AROS DE AGILIDAD|BANDA DE CAPITAN|CUBITT|UNIFORME FUTBOL|^MEDIAS? '), 'T', 'Fútbol', ''),
  # ---- Tienda: ropa, uniformes y entrenamiento ---------------------------
  (R(r'T[ -]?SHIRT|POLOSHIRT|PANTALON|FALDA|FALDITA|^TOP |BLUSA|VESTIDO|LICRA|SUERA|SHORT|^MANGAS$|CALCETINES|TOALLA|RODILLERA|COLCHONETA|RESISTANGE'), 'T', 'Ropa y accesorios', ''),
  (R(r'BOLSA DE REGALO'), 'T', 'Ropa y accesorios', ''),

  # ---- Cafetería: insumos (no se venden) ---------------------------------
  (R(r'QUESO MOZZARELLA|MANTECA VEGETAL|^SALSA BBQ|CATCHUP|PAN MEDIA|REDROT|MIEL RINCON|^LECHE$|^SALCHICHA$|PORCION DE SALAMI|QUESO BLANCO DE FREIR|^LONGANIZA$'), 'C', 'Insumos de cocina', 'Insumo'),
  (R(r'^PIZZA|SLIDE DE PIZZA|INGREDIENTES ADICIONALES PIZZA'), 'C', 'Pizzas', 'Preparado'),
  # ---- Cafetería: bar ----------------------------------------------------
  (R(r'^TRAGO|BEBIDA APEROL|^FROZEN|DAIKIRI|SANGRIA|^PALOMA$'), 'C', 'Tragos y cócteles', 'Preparado'),
  (R(r'SMIRNOFF ICE|WHITE CLAW|BUZZBALL|TRIPP?ING? ANIMAL'), 'C', 'Tragos y cócteles', 'Reventa'),
  (R(r'DESCORCHE|VASO CON HIELO'), 'C', 'Servicios de bar', 'Servicio'),
  (R(r'FUNDA DE HIELO'), 'C', 'Servicios de bar', 'Reventa'),
  (R(r'ESPUMANTE|CAVA BRUT'), 'C', 'Licores y vinos', 'Reventa'),
  (R(r'CERVEZA|MICHELOB|MILLER|COORS|HEINEKEN|PRESIDENTE|CORONA|CORONITA|STELLA|PERONI|MODELO|PAULANER|ERDINGER|^5,0|SAPPORO|BLUE MOON|GROLSCH|DAMM|ALHAMDRA|DESPERADO|KELER'), 'C', 'Cervezas', 'Reventa'),
  (R(r'FERNET|TEQUILA|SIBONEY|GIBSON|STOLI|BEEFEATER|WHITLEY|JIMADOR|MACORIX|LAGRANGE|OJO DE TIGRE|BOTELLA +APEROL|BRUGAL|WHISKY|CHIVA|RESERVA BOTELLA|RON 1888 BOTELLA|EXTRA VIEJO BOTELLA'), 'C', 'Licores y vinos', 'Reventa'),
  # ---- Cafetería: café ---------------------------------------------------
  (R(r'NESCAFE'), 'C', 'Café', 'Preparado'),
  (R(r'FRAPPUC|CHOCOLATE FRIO'), 'C', 'Café', 'Reventa'),
  # ---- Cafetería: bebidas sin alcohol ------------------------------------
  (R(r'GATORADE|GATORLIT|GATORLYTE|SUEROX|RED BULL|MONSTER|C4 EN|VITARAI'), 'C', 'Deportivas y energizantes', 'Reventa'),
  (R(r'^AGUA (DASANI|EVIAN|FIJI|PANNA|PUREZA|PLANETA)|ICELANDIC|^PERRIER|^SAN PELLEGRINO 250'), 'C', 'Aguas', 'Reventa'),
  (R(r'^JUGO|RICA|WELCH.S (ORANGE|GRAPE|FRUIT)|FRUIT PARADISE|CLAMATO|V8 SPLASH|OCEAN SPRAY|AGUA DE COCO|ENTEREX|COMPOTA'), 'C', 'Jugos y lácteos', 'Reventa'),
  (R(r'COCA.COLA|SPRITE|COUNTRY CLUB|CANADA|RED ROCK|7 UP|TOP TOP|MOUNTAIN DEW|JARRITOS|TONICA|SPAR[KL]|LA CROIX|SAN PEREGRINO|REFRESCO'), 'C', 'Refrescos', 'Reventa'),
  # ---- Cafetería: helados ------------------------------------------------
  (R(r'HELADO|PALETA|MAGNUM|^COPA|ICE POPS|BLUE RIBBON|MORDISKO|SANDWICH BON'), 'C', 'Helados y paletas', 'Reventa'),
  # ---- Cafetería: comida -------------------------------------------------
  (R(r'^PASTA|TOSTADA|^SANDWICH|^HAMBURG|HOT DOG|^WRAP|^SERVICIO|ALITAS|PORCION DE ALAS|PAPAS FRITAS|YUQUITAS FRITAS|^PLATO|ARROZ FRITO|CROISSANT JAMON|NACHOS MARIA|PECHUGA|CROQUETAS'), 'C', 'Comida', 'Preparado'),
  (R(r'EMPANAD|QUIPE|^PORCION|MAC & CHEESE|SOPA ISIM|CASABE'), 'C', 'Comida', 'Reventa'),
  # ---- Cafetería: dulces y snacks ----------------------------------------
  (R(r'GALLETA|OREO|CHOKIS|EMPERADOR|FLORENTINA|MAMUT|MISSCOOKIE|TATES|KIRKLAND CHOCOLATE|NATURE VALL|KELLOGG|NUTRI GRAIN|BARRA CEREAL|RIP VAN|MY MOTTO|CRACHI|CROSTATA|CINNAMON|MUFFIN|BIZCOCHO|BROWNIES|DONAS|CAPACILLOS|CROISSANT|MINI ANGELITOS|DINO |PIRULIN|TAW|CRACKETS|GOLDFISH|JOLIE|FLIPS'), 'C', 'Galletas y repostería', 'Reventa'),
  (R(r'M&M|SNICKERS|TWIX|MILKYWAY|KITKAT|HERSHEY|KINDER|SKITTLES|STARBURST|TIC TAC|MENTOS|AIR HEADS|RING POP|PUSHPOP|BABY BOTTLE|LENGUITAS|GOMITAS|FRUIT SNACKS|NUCITA|CHOCOLATE|CRICRI|DULCE DE LECHE NATULAC|COQUITO'), 'C', 'Chocolates y dulces', 'Reventa'),
  (R(r'LAYS|DORITOS|CHEETOS|PRINGLES|RUFFLES|NATUCHIPS|ZAMBOS|RANCH|RANC\b|TAQUE|ZIBAS|YUMMI|HOJUELITA|CARLES|CHICHARRON|PLATANITOS|TAKIS|DETODITO|MANI|PISTACHIO|CARIBAS|PALOMITAS|CORNETAS|CHEMILO|RABITOS|TOSTONES'), 'C', 'Snacks', 'Reventa'),
]


def classify(name):
    n = name.upper()
    for pat, dest, grp, extra in RULES:
        if pat.search(n):
            return dest, grp, extra
    raise SystemExit(f'Sin clasificar: {name!r} — agregar una regla')


# --------------------------------------------------------------------------
# Nombres de la cafetería: solo errores de tipeo evidentes. El original queda
# en la columna "nombre original".
# --------------------------------------------------------------------------
NAME_FIXES = [
    (r'FRAPPUCINNO', 'FRAPPUCCINO'),
    (r'JOHONNY|JHONNY|JOHNINIE', 'JOHNNIE'),
    (r'\bCHIVA 18', 'CHIVAS 18'),
    (r'DAIKIRI', 'DAIQUIRI'),
    (r'(?i)^sparlink ice back berry$', 'SPARKLING ICE BLACKBERRY'),
    (r'(?i)^sparlink|^SPARKLINK', 'SPARKLING'),
    (r'PINNEAPPLE', 'PINEAPPLE'),
    (r'ENREGY', 'ENERGY'),
    (r'RAIMBOW', 'RAINBOW'),
    (r'FROZZEN', 'FROZEN'),
    (r'CKASICA', 'CLASICA'),
    (r'GATORLIT REVER', 'GATORLIT RECOVER'),
    (r'FRESA KIWIT', 'FRESA KIWI'),
    (r'BELGION', 'BELGIAN'),
    (r'WEIBBIER', 'WEISSBIER'),
    (r'CERVEZA ENDINGER', 'CERVEZA ERDINGER'),
    (r'ALHAMDRA RESERVAS', 'ALHAMBRA RESERVA'),
    (r'^WHERE CERVEZA', 'CERVEZA'),
    (r'LARGER BEER', 'LAGER BEER'),
    (r',S\b', "'S"),
    (r'PREMIUN', 'PREMIUM'),
    (r'’', "'"),
    (r'SAN PEREGRINO|SANPELLEGRINO', 'SAN PELLEGRINO'),
    (r'PANNA PEQUENA', 'PANNA PEQUEÑA'),
    (r'^128PALETA DAZS CHOCOLATE CHOC$', 'PALETA DAZS CHOCOLATE'),
    (r'\bDANS\b', 'DAZS'),
    (r'STRAWBERRIRS', 'STRAWBERRIES'),
    (r'NATURE VALLERY', 'NATURE VALLEY'),
    (r'CATCHUP', 'KETCHUP'),
    (r'CHIP ALOY', 'CHIPS AHOY'),
    (r'^HAMBURGUER$', 'HAMBURGUESA'),
    (r'^TAQUE CHILETOR', 'TAQUERITOS CHILE TOREADO'),
    (r'^TAQUER CHILETOR', 'TAQUERITOS CHILE TOREADO'),
    (r'^VITARAI ', 'VITARAIN '),
    (r'CARDE ASADA', 'CARNE ASADA'),
    (r'CHEEDAR', 'CHEDDAR'),
    (r'BANGUETTE', 'BAGUETTE'),
    (r'^ROM MACORIX', 'RON MACORIX'),
    (r'CLUD CRACKERS', 'CLUB CRACKERS'),
    (r'FRUIT PUNCH 350M$', 'FRUIT PUNCH 350ML'),
    (r' - ', ' '),
    (r'LATA- ', 'LATA '),
]


def clean_name(name):
    n = name
    for pat, rep in NAME_FIXES:
        n = re.sub(pat, rep, n)
    n = re.sub(r'\s+', ' ', n).strip().upper()
    return n


# Posibles duplicados (mismo producto con dos códigos). Se cargan los dos;
# el dueño decide cuál desactivar.
DUPES = [
    ('110000100619', '4002103248248'),                 # ERDINGER
    ('110000100303', '790330021249'),                  # JUGO RICA PERA
    ('110000100302', '014100077602'),                  # GOLDFISH
    ('110000100294', '110000100297'),                  # GALLETAS COCO MARTIN PEQ.
    ('1140', '1250'),                                  # COPA CHOCOLATE PRISCILLA
    ('110000100560', '110000100654'),                  # PALOMA
    ('110000100585', '7460855234990'),                 # EXTRA VIEJO
    ('110000100334', '110000100499', '1034'),          # PAPAS FRITAS
    ('110000100335', '110000100500'),                  # YUQUITAS
    ('154545845', '750894612550'),                     # TAQUERITOS CHILE TOREADO
    ('1170', '750894614301'),                          # RANCH BUFFE RANCH
    ('110000100612', '653981779023'),                  # GALLETAS DE AVENA
    ('110000100331', '1216'),                          # HAMBURGUESA
]

EXTRA_NOTES = {
    '1219': 'Es una OFERTA (3x2), no un producto: mejor configurarla en Ofertas sobre MICHELOB ULTRA 330 ML.',
    '1181': 'Mejor como modificador de las pizzas que como producto suelto.',
    '110000100662': 'Nombre genérico: ¿qué frozen es?',
    '1653265': 'Código roto (7 dígitos): queda solo como SKU.',
    '9661931395': 'Código roto (10 dígitos): queda solo como SKU.',
    '154545845': 'Código roto (9 dígitos): queda solo como SKU.',
    '1176': 'Costo 4,278.35 por unidad: ¿es un bloque/caja? Definir la unidad antes de cargar existencia.',
}

CAT_ORDER = [
    'Cervezas', 'Licores y vinos', 'Tragos y cócteles', 'Refrescos', 'Aguas',
    'Jugos y lácteos', 'Deportivas y energizantes', 'Café', 'Pizzas', 'Comida',
    'Snacks', 'Galletas y repostería', 'Chocolates y dulces',
    'Helados y paletas', 'Servicios de bar', 'Insumos de cocina',
]
TIPO_LABEL = {
    'Reventa': 'Reventa (lleva stock)',
    'Preparado': 'Preparado (sin stock propio)',
    'Servicio': 'Servicio (sin stock)',
    'Insumo': 'Insumo (no se vende)',
}


def build():
    rows = load()
    total_rows = len(rows)
    total_cost = round(sum(r['total'] for r in rows), 2)

    dupe_of = {}
    for group in DUPES:
        for c in group:
            dupe_of[c] = [o for o in group if o != c]
    by_code = {r['code']: r for r in rows}
    for group in DUPES:
        for c in group:
            assert c in by_code, f'DUPES: código {c} no existe'

    caf, tienda, decidir = [], [], []
    for r in rows:
        dest, grp, extra = classify(r['name'])
        r['group'] = grp
        if dest == 'T':
            tienda.append(r)
            continue
        if dest == 'X':
            r['reason'] = extra
            decidir.append(r)
            continue

        tipo = extra
        if r['code'] == '1219':
            tipo = 'Oferta'
        r['tipo'] = tipo
        r['clean'] = clean_name(r['name'])
        r['barcode'] = barcode_for(r['code'])

        notes = []
        cost = r['cost']
        if abs(cost - 1.0) < 1e-9:
            notes.append('Costo 1.00 del sistema viejo era relleno: se deja sin costo.')
            cost = None
        elif cost == 0:
            cost = None
            if r['stock'] > 0:
                notes.append(f"Tiene {r['stock']:g} en existencia pero SIN costo.")
        r['cost_out'] = cost

        stock = r['stock']
        if tipo in ('Reventa', 'Insumo'):
            r['stock_out'] = max(0.0, stock)
            if stock < 0:
                notes.append(f'Existencia negativa ({stock:g}) en el sistema viejo: se carga 0.')
        else:
            r['stock_out'] = None
            if stock > 0:
                notes.append(f'Es {tipo.lower()} pero el sistema viejo le tenía {stock:g} en existencia.')
        if r['code'] in dupe_of:
            names = ', '.join(f"{by_code[o]['name']} ({o})" for o in dupe_of[r['code']])
            notes.append(f'Posible duplicado de: {names}.')
        if r['code'] in EXTRA_NOTES:
            notes.append(EXTRA_NOTES[r['code']])
        if r['clean'] != re.sub(r'\s+', ' ', r['name']).strip().upper():
            r['renamed'] = True
        r['notes'] = notes
        caf.append(r)

    # ---- Cuadre: nada se pierde en la división ---------------------------
    assert len(caf) + len(tienda) + len(decidir) == total_rows
    split_cost = round(sum(r['total'] for r in caf + tienda + decidir), 2)
    assert split_cost == total_cost, (split_cost, total_cost)
    assert len({r['clean'] for r in caf}) == len(caf), 'Nombres repetidos en cafetería'

    key = lambda r: (CAT_ORDER.index(r['group']), r['clean'])
    caf.sort(key=key)
    productos = [r for r in caf if r['tipo'] != 'Insumo']
    insumos = [r for r in caf if r['tipo'] == 'Insumo']

    write_cafeteria(productos, insumos, tienda, decidir, total_rows, total_cost)
    write_tienda(tienda)
    write_decidir(decidir)
    write_diagnostico(productos + insumos)

    print(f'{total_rows} renglones · RD${total_cost:,.2f} a costo')
    for label, part in (('Cafetería', caf), ('Tienda', tienda), ('Por decidir', decidir)):
        print(f'  {label:<12} {len(part):>4}  RD${sum(r["total"] for r in part):>12,.2f}')
    print(f'  Cafetería = {len(productos)} productos + {len(insumos)} insumos')


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
BOLD = Font(bold=True)
HEAD_FILL = PatternFill('solid', fgColor='DDE3EA')
PRICE_FILL = PatternFill('solid', fgColor='FFF2A8')
MONEY = '#,##0.00'
QTY = '#,##0'


def sheet(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'



def save(wb, path):
    """Guarda y deja las rutas internas RELATIVAS.

    openpyxl escribe `Target="/xl/worksheets/sheet1.xml"` (absoluto) y el
    paquete `excel` de la app busca `xl/` + target → no encuentra la hoja y
    el importador truena ("Null check operator"). Excel acepta las dos formas.
    """
    wb.save(path)
    rels = 'xl/_rels/workbook.xml.rels'
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == rels:
                data = data.replace(b'Target="/xl/', b'Target="')
            z.writestr(info, data)


def put_text(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value if value else None)
    c.number_format = '@'


def write_cafeteria(productos, insumos, tienda, decidir, total_rows, total_cost):
    wb = Workbook()

    # Hoja 1 = lo que lee Productos → Importar catálogo (solo la PRIMERA hoja,
    # encabezado en la fila 1). Encabezados reconocidos por CatalogCsvParser:
    # nombre, precio, costo, codigo (→ SKU), codigo de barras, categoria, activo.
    # Las demás columnas el importador las ignora.
    ws = wb.active
    ws.title = 'Productos'
    sheet(ws,
          ['nombre', 'precio', 'costo', 'codigo', 'codigo de barras', 'categoria',
           'activo', 'tipo', 'existencia a cargar', 'existencia sistema viejo',
           'notas', 'nombre original'],
          [42, 12, 11, 16, 16, 24, 8, 26, 12, 12, 70, 42])
    for i, r in enumerate(productos, start=2):
        ws.cell(row=i, column=1, value=r['clean'])
        ws.cell(row=i, column=2).fill = PRICE_FILL
        ws.cell(row=i, column=2).number_format = MONEY
        c = ws.cell(row=i, column=3, value=r['cost_out'])
        c.number_format = MONEY
        put_text(ws, i, 4, r['code'])
        put_text(ws, i, 5, r['barcode'])
        ws.cell(row=i, column=6, value=r['group'])
        ws.cell(row=i, column=7, value='Si')
        ws.cell(row=i, column=8, value=TIPO_LABEL.get(r['tipo'], r['tipo']))
        c = ws.cell(row=i, column=9, value=r['stock_out'])
        c.number_format = QTY
        c = ws.cell(row=i, column=10, value=r['stock'])
        c.number_format = QTY
        ws.cell(row=i, column=11, value=' '.join(r['notes']) or None)
        ws.cell(row=i, column=12, value=r['name'] if r.get('renamed') else None)

    # Insumos: no son productos de venta; NO van por el importador de catálogo.
    wi = wb.create_sheet('Insumos')
    sheet(wi, ['nombre', 'codigo', 'costo', 'existencia a cargar',
               'existencia sistema viejo', 'notas', 'nombre original'],
          [36, 14, 11, 12, 12, 60, 36])
    for i, r in enumerate(insumos, start=2):
        wi.cell(row=i, column=1, value=r['clean'])
        put_text(wi, i, 2, r['code'])
        wi.cell(row=i, column=3, value=r['cost_out']).number_format = MONEY
        wi.cell(row=i, column=4, value=r['stock_out']).number_format = QTY
        wi.cell(row=i, column=5, value=r['stock']).number_format = QTY
        wi.cell(row=i, column=6, value=' '.join(r['notes']) or None)
        wi.cell(row=i, column=7, value=r['name'] if r.get('renamed') else None)

    # Revisar: una fila por nota, agrupadas por tipo de problema.
    wr = wb.create_sheet('Revisar')
    sheet(wr, ['problema', 'nombre', 'codigo', 'categoria', 'detalle'],
          [30, 42, 16, 24, 80])
    kinds = [
        ('Posible duplicado', 'Posible duplicado'),
        ('Sin costo con existencia', 'SIN costo'),
        ('Costo de relleno (1.00)', 'Costo 1.00'),
        ('Otros', None),
    ]
    out = []
    for r in productos + insumos:
        for n in r['notes']:
            if n.startswith('Existencia negativa'):
                continue  # van resumidas abajo, son muchas
            kind = next((k for k, pref in kinds if pref and pref in n), 'Otros')
            out.append((kind, r, n))
    out.sort(key=lambda t: ([k for k, _ in kinds].index(t[0]), t[1]['clean']))
    row = 2
    for kind, r, n in out:
        wr.cell(row=row, column=1, value=kind)
        wr.cell(row=row, column=2, value=r['clean'])
        put_text(wr, row, 3, r['code'])
        wr.cell(row=row, column=4, value=r['group'])
        wr.cell(row=row, column=5, value=n)
        row += 1
    negs = [r for r in productos + insumos if r['stock'] < 0 and r['tipo'] in ('Reventa', 'Insumo')]
    wr.cell(row=row + 1, column=1,
            value=f'Además: {len(negs)} productos de reventa con existencia NEGATIVA '
                  f'en el sistema viejo se cargan en 0 (ver columna "notas" en Productos).').font = BOLD

    # Resumen
    ws2 = wb.create_sheet('Resumen')
    sheet(ws2, ['categoria', 'productos', 'reventa', 'preparados/servicios',
                'con existencia', 'unidades a cargar', 'valor a costo'],
          [26, 10, 10, 14, 12, 14, 16])
    ws2.freeze_panes = None
    stats = collections.OrderedDict((c, [0, 0, 0, 0, 0.0, 0.0]) for c in CAT_ORDER)
    for r in productos + insumos:
        s = stats[r['group']]
        s[0] += 1
        if r['tipo'] in ('Reventa', 'Insumo', 'Oferta'):
            s[1] += 1
        else:
            s[2] += 1
        if r['stock_out']:
            s[3] += 1
            s[4] += r['stock_out']
            s[5] = round(s[5] + r["stock_out"] * (r["cost_out"] or 0), 2)
    row = 2
    for cat, s in stats.items():
        ws2.append([cat, *s])
        ws2.cell(row=row, column=6).number_format = QTY
        ws2.cell(row=row, column=7).number_format = MONEY
        row += 1
    ws2.append(['TOTAL CAFETERÍA'] + [sum(s[i] for s in stats.values()) for i in range(6)])
    for col in range(1, 8):
        ws2.cell(row=row, column=col).font = BOLD
    ws2.cell(row=row, column=6).number_format = QTY
    ws2.cell(row=row, column=7).number_format = MONEY
    row += 2
    notes = [
        f'Fuente: ARTICULOS DE INVENTARIO.csv, {total_rows} renglones, RD${total_cost:,.2f} a costo (con negativos).',
        f'Cafetería: {len(productos)} productos + {len(insumos)} insumos · Tienda: {len(tienda)} · Por decidir: {len(decidir)}.',
        'El archivo viejo NO trae precio de venta: la columna PRECIO (amarilla) hay que llenarla antes de importar.',
        'Importar: Productos → Importar catálogo → este Excel (lee solo la hoja "Productos").',
        'En el importador elegir el ITBIS correcto en "Impuesto por defecto" (por defecto propone EXENTO si existe).',
        'El importador NO carga existencias: la columna "existencia a cargar" va en un segundo paso.',
        'Valor a costo = existencia a cargar × costo; los productos sin costo suman 0.',
    ]
    for n in notes:
        ws2.cell(row=row, column=1, value=n)
        row += 1

    save(wb, os.path.join(OUT, 'CAFETERIA_Estancia_Nueva.xlsx'))


def write_raw(path, title, rows, extra_header, extra_value):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    sheet(ws, ['codigo', 'descripcion', 'ultimo costo', 'existencia', 'costo total', extra_header],
          [16, 44, 12, 11, 13, 40])
    for i, r in enumerate(rows, start=2):
        put_text(ws, i, 1, r['code'])
        ws.cell(row=i, column=2, value=r['name'])
        ws.cell(row=i, column=3, value=r['cost']).number_format = MONEY
        ws.cell(row=i, column=4, value=r['stock']).number_format = QTY
        ws.cell(row=i, column=5, value=r['total']).number_format = MONEY
        ws.cell(row=i, column=6, value=extra_value(r))
    n = len(rows) + 2
    ws.cell(row=n, column=2, value=f'TOTAL ({len(rows)} artículos)').font = BOLD
    c = ws.cell(row=n, column=5, value=round(sum(r['total'] for r in rows), 2))
    c.number_format = MONEY
    c.font = BOLD
    save(wb, path)


def write_tienda(rows):
    order = ['Pádel', 'Fútbol', 'Ropa y accesorios', 'Servicios del club']
    rows = sorted(rows, key=lambda r: (order.index(r['group']), r['name']))
    write_raw(os.path.join(OUT, 'TIENDA_Estancia_Nueva.xlsx'), 'Tienda', rows,
              'grupo', lambda r: r['group'])


def write_decidir(rows):
    rows = sorted(rows, key=lambda r: (r['group'], r['name']))
    write_raw(os.path.join(OUT, 'POR_DECIDIR_Estancia_Nueva.xlsx'), 'Por decidir', rows,
              'por qué', lambda r: f"{r['group']}: {r['reason']}")


# --------------------------------------------------------------------------
# Diagnóstico (solo lectura) del negocio de la cafetería.
# UNA sola consulta: el SQL Editor de Supabase solo muestra la última.
# Columnas que no están en supabase/schema.sql (la BD viva diverge) se leen
# con to_jsonb(fila) ->> 'col' para que la consulta nunca truene.
# --------------------------------------------------------------------------
def q(v):
    return "'" + str(v).replace("'", "''") + "'"


DIAG_TPL = """-- ============================================================================
-- DIAGNÓSTICO — Cafetería Estancia Nueva Sports Club · business __BID__
--
-- ⚠ ARCHIVO GENERADO por build_estancia_nueva.py. Editar el .py, no esto.
--
-- SOLO LEE. Correr completo y pegarme el resultado. Es UNA sola consulta a
-- propósito (el SQL Editor de Supabase solo muestra la última). Sale una
-- tabla (seccion, detalle):
--   1) el negocio y los otros negocios del mismo dueño
--   2) ajustes: cocina, inventario, ITBIS, menús, bodegas, áreas de comanda
--   3) catálogo actual en números
--   4) cruce con los __N__ artículos del archivo (por código y por nombre)
--   5) los que YA existen: con su precio en el sistema
--   6) productos del sistema que NO están en el archivo
--   7) insumos que ya existen con un código/nombre del archivo
--   8) categorías actuales
-- ============================================================================
with
biz as (
  select '__BID__'::uuid as id
),
archivo(codigo, nombre, original, categoria, tipo) as (
  values
__VALUES__
),
n as (  -- nombres normalizados: sin acentos, mayúsculas, espacios simples
  select a.*,
         upper(regexp_replace(translate(btrim(a.nombre),   'áéíóúÁÉÍÓÚñÑ', 'aeiouAEIOUnN'), '\\s+', ' ', 'g')) as n1,
         upper(regexp_replace(translate(btrim(a.original), 'áéíóúÁÉÍÓÚñÑ', 'aeiouAEIOUnN'), '\\s+', ' ', 'g')) as n2
  from archivo a
),
mi as (
  select m.id, m.name, m.price, m.cost, m.sku, m.barcode, m.is_active, m.created_at,
         c.name as cat,
         upper(regexp_replace(translate(btrim(m.name), 'áéíóúÁÉÍÓÚñÑ', 'aeiouAEIOUnN'), '\\s+', ' ', 'g')) as nn,
         coalesce((to_jsonb(m) ->> 'is_inventory_tracked')::boolean, false) as tracked,
         exists (select 1 from public.order_items oi where oi.product_id = m.id) as vendido
  from public.menu_items m
  left join public.categories c on c.id = m.category_id
  where m.business_id = (select id from biz)
),
cruce as (
  select n.*, x.id as mid, x.name as mname, x.price, x.cat, x.vendido, x.is_active as mactivo, x.via
  from n
  left join lateral (
    select m.*, case when m.sku = n.codigo or m.barcode = n.codigo then 'código' else 'nombre' end as via
    from mi m
    where m.sku = n.codigo or m.barcode = n.codigo or m.nn in (n.n1, n.n2)
    order by (m.sku = n.codigo or m.barcode = n.codigo) desc, m.is_active desc, m.created_at
    limit 1
  ) x on true
),
ii as (
  select i.id, i.name, i.sku, to_jsonb(i) ->> 'barcode' as barcode,
         upper(regexp_replace(translate(btrim(i.name), 'áéíóúÁÉÍÓÚñÑ', 'aeiouAEIOUnN'), '\\s+', ' ', 'g')) as nn
  from public.inventory_items i
  where i.business_id = (select id from biz)
),
filas(orden, sub, seccion, detalle) as (
  -- 1) Negocio
  select 10, 0, '1. Negocio',
         b.business_name || coalesce(' · ' || b.branch_name, '') || ' · tipo ' || coalesce(b.business_type, '—')
         || ' · ' || b.status || ' · ' || b.domain || ' · creado ' || to_char(b.created_at, 'DD/MM/YYYY')
  from public.businesses b where b.id = (select id from biz)
  union all
  select 10, 0, '1. Negocio', 'NO EXISTE un negocio con ese id'
  where not exists (select 1 from public.businesses b where b.id = (select id from biz))
  union all
  select 11, 0, '1. Otro negocio del mismo dueño',
         o.business_name || coalesce(' · ' || o.branch_name, '') || ' · ' || o.id || ' · ' || o.status
         || ' · ' || (select count(*) from public.menu_items m where m.business_id = o.id) || ' productos'
  from public.businesses o
  where o.owner_id = (select owner_id from public.businesses where id = (select id from biz))
    and o.id <> (select id from biz)

  -- 2) Ajustes
  union all
  select 20, 0, '2. Ajustes',
         'kitchen_enabled = ' || coalesce(to_jsonb(bs) ->> 'kitchen_enabled', '¿?')
         || ' · inventory_mode = ' || coalesce(to_jsonb(bs) ->> 'inventory_mode', '¿?')
         || ' · allow_negative_stock = ' || coalesce(to_jsonb(bs) ->> 'allow_negative_stock', '¿?')
         || ' · service_fee_enabled = ' || coalesce(to_jsonb(bs) ->> 'service_fee_enabled', '¿?')
         || ' · printerless_kitchen = ' || coalesce(to_jsonb(bs) ->> 'printerless_kitchen', '¿?')
         || ' · auto_print_order = ' || coalesce(to_jsonb(bs) ->> 'auto_print_order', '¿?')
         || ' · moneda = ' || coalesce(to_jsonb(bs) ->> 'currency', to_jsonb(bs) ->> 'currency_code', '¿?')
  from public.business_settings bs
  where bs.business_id = (select id from biz)
  union all
  select 20, 0, '2. Ajustes', 'SIN fila en business_settings'
  where not exists (select 1 from public.business_settings bs where bs.business_id = (select id from biz))
  union all
  select 21, 0, '2. Ajustes', 'Impuestos: ' || coalesce((
           select string_agg(t.name || ' ' || t.rate || '%'
                  || case when coalesce(t.is_active, true) then '' else ' (inactivo)' end
                  || case when coalesce((to_jsonb(t) ->> 'is_service_fee')::boolean, false) then ' (is_service_fee!)' else '' end,
                  ', ' order by t.rate)
           from public.taxes t where t.business_id = (select id from biz)), 'NINGUNO')
  union all
  select 22, 0, '2. Ajustes', 'Menús activos: ' || coalesce((
           select string_agg(me.name, ', ') from public.menus me
           where me.business_id = (select id from biz) and coalesce(me.is_active, true)), 'NINGUNO')
  union all
  select 23, 0, '2. Ajustes', 'Bodegas: ' || coalesce((
           select string_agg(w.name
                  || case when w.is_main then ' (principal)' else '' end
                  || case when coalesce(w.is_active, true) then '' else ' (INACTIVA)' end, ', '
                  order by w.is_main desc, w.created_at asc nulls first, w.id asc)
           from public.warehouses w where w.business_id = (select id from biz)), 'NINGUNA')
         || ' · la venta descuenta de: ' || coalesce((
           select w.name from public.warehouses w
           where w.business_id = (select id from biz)
           order by w.is_main desc, w.created_at asc nulls first, w.id asc limit 1), '—')
  union all
  select 24, 0, '2. Ajustes', 'Áreas de comanda: ' || coalesce((
           select string_agg(a.name || ' (' || a.code || ')', ', ' order by a.name)
           from public.print_areas a
           where a.business_id = (select id from biz) and a.is_active
             and a.code not in ('cashier', 'fiscal', 'cash_close')), 'ninguna')

  -- 3) Catálogo actual
  union all
  select 30, 0, '3. Catálogo actual',
         (select count(*) from mi) || ' productos (' || (select count(*) from mi where is_active) || ' activos, '
         || (select count(*) from mi where price > 0) || ' con precio, '
         || (select count(*) from mi where vendido) || ' con ventas, '
         || (select count(*) from mi where tracked) || ' inventariables) · '
         || (select count(*) from public.categories c where c.business_id = (select id from biz)) || ' categorías · '
         || (select count(*) from ii) || ' insumos'

  -- 4) Cruce con el archivo
  union all
  select 40, 0, '4. Cruce con el archivo',
         count(*) filter (where via = 'código') || ' ya existen por CÓDIGO · '
         || count(*) filter (where via = 'nombre') || ' por NOMBRE · '
         || count(*) filter (where mid is null) || ' no existen · de ' || count(*) || ' del archivo ('
         || count(*) filter (where tipo = 'Insumo') || ' son insumos)'
  from cruce

  -- 5) Los que ya existen, con su precio
  union all
  select 50, 0, '5. Ya existe · ' || c.categoria,
         c.codigo || ' · ' || c.nombre || case when c.mname <> c.nombre then ' ↔ ' || c.mname else '' end
         || ' · $' || c.price || ' · por ' || c.via || ' · ' || coalesce(c.cat, 'SIN CATEGORÍA')
         || case when c.vendido then ' · CON ventas' else '' end
         || case when c.mactivo then '' else ' · INACTIVO' end
  from cruce c where c.mid is not null

  -- 6) Productos del sistema que no están en el archivo
  union all
  select 60, 0, '6. Está en el sistema, NO en el archivo',
         m.name || ' · $' || m.price || ' · sku ' || coalesce(m.sku, '—') || ' · ' || coalesce(m.cat, 'SIN CATEGORÍA')
         || case when m.vendido then ' · CON ventas' else '' end
         || case when m.is_active then '' else ' · INACTIVO' end
  from mi m
  where not exists (select 1 from cruce c where c.mid = m.id)

  -- 7) Insumos que ya existen
  union all
  select 70, 0, '7. Insumo que ya existe',
         i.name || ' · sku ' || coalesce(i.sku, '—') || ' · barcode ' || coalesce(i.barcode, '—')
         || ' · movimientos ' || (select count(*) from public.inventory_movements mv where mv.item_id = i.id)
  from ii i
  where exists (select 1 from n where n.codigo in (i.sku, i.barcode) or i.nn in (n.n1, n.n2))

  -- 8) Categorías
  union all
  select 80, c.position, '8. Categoría',
         c.name || ' · posición ' || c.position
         || ' · ' || (select count(*) from public.menu_items m where m.category_id = c.id) || ' productos'
         || case when c.is_active then '' else ' · INACTIVA' end
  from public.categories c
  where c.business_id = (select id from biz)
)
select seccion, detalle
from filas
order by orden, sub, seccion, detalle;
"""


def write_diagnostico(rows):
    values = ',\n'.join(
        '    ({}, {}, {}, {}, {})'.format(
            q(r['code']), q(r['clean']), q(r['name']), q(r['group']), q(r['tipo']))
        for r in rows)
    sql = (DIAG_TPL.replace('__BID__', BID)
           .replace('__N__', str(len(rows)))
           .replace('__VALUES__', values))
    with open(os.path.join(OUT, f'00_diagnostico_{BID[:8]}.sql'), 'w') as f:
        f.write(sql)


if __name__ == '__main__':
    build()
