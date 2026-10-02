#!/usr/bin/env python3
"""Genera la carga del catálogo de la CAFETERÍA de Estancia Nueva (business 85924083).

Uso:
    <python con openpyxl> build_import_cafeteria.py [--precios ARCHIVO] [--out CARPETA]

Toma los 636 productos y 12 insumos de la cafetería tal como los clasifica
build_estancia_nueva.prepare() y escribe, desde las plantillas _tpl_*.sql:

    IMPORT_CAFETERIA.sql      la carga, una sola transacción (re-ejecutable)
    99_rollback_cafeteria.sql la deshace

--precios: CSV o Excel (primera hoja) con columnas `codigo` y `precio` (ITBIS
incluido). Sirve el mismo CAFETERIA_Estancia_Nueva.xlsx con la columna amarilla
llena. Sin precios, todo entra INACTIVO en RD$0; con precios se regenera y se
vuelve a correr: los que traen precio se activan.

--proisa: la lista de precios del sistema viejo (ProIsa_lista.csv: codigo,
descripcion, ..., precio_rd). Ver leer_proisa(): ProIsa guarda el precio SIN
ITBIS. Escribe además precios_cafeteria_revision.csv (uno por producto, con la
regla que se aplicó) para revisarlo con el dueño.

Editar ESTE archivo o las plantillas y regenerar; los .sql generados no se tocan
a mano.
"""

import argparse
import csv
import re
import sys
import unicodedata
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from build_estancia_nueva import BID, CAT_ORDER, barcode_for, prepare

HERE = Path(__file__).resolve().parent

CATEGORIAS = [c for c in CAT_ORDER if c != 'Insumos de cocina']
BEBIDAS = {'Cervezas', 'Licores y vinos', 'Tragos y cócteles', 'Refrescos',
           'Aguas', 'Jugos y lácteos', 'Deportivas y energizantes', 'Café'}
COCINA = {'Comida', 'Pizzas'}


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


def precio_de(v):
    """'1,250.00' / 1250 / '' → float o None. 0 cuenta como sin precio."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        x = float(v)
    else:
        s = str(v).strip().replace('RD$', '').replace('$', '').replace(',', '')
        if not s:
            return None
        try:
            x = float(s)
        except ValueError:
            raise SystemExit(f'Precio que no es número: {v!r}')
    if x < 0:
        raise SystemExit(f'Precio negativo: {v!r}')
    return round(x, 2) if x > 0 else None


def leer_precios(path):
    path = Path(path).expanduser()
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
        filas = ws.iter_rows(values_only=True)
        hdr = [str(h or '').strip().lower() for h in next(filas)]
        rows = [dict(zip(hdr, r)) for r in filas]
    else:
        with open(path, encoding='utf-8-sig', newline='') as fh:
            rows = [{(k or '').strip().lower(): v for k, v in r.items()}
                    for r in csv.DictReader(fh)]
    if not rows or 'codigo' not in rows[0] or 'precio' not in rows[0]:
        raise SystemExit(f'{path.name}: hacen falta las columnas "codigo" y "precio".')
    precios = {}
    for r in rows:
        code = str(r['codigo'] or '').strip()
        if not code:
            continue
        if code.endswith('.0'):          # Excel guardó el código como número
            code = code[:-2]
        p = precio_de(r['precio'])
        if code in precios and precios[code] != p:
            raise SystemExit(f'{path.name}: el código {code} trae dos precios distintos.')
        precios[code] = p
    return precios, path.name


def _norm(s):
    s = unicodedata.normalize('NFKD', str(s).upper()).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip()


def _redondo(x, paso=5):
    return abs(x - round(x / paso) * paso) < 0.06


# Cuántos precios de ProIsa ya son el precio FINAL. Si la lista cambia y el
# número se mueve, el generador se detiene para que alguien lo mire.
PROISA_YA_FINALES = 24


def leer_proisa(path, productos, conocidos, esperados=PROISA_YA_FINALES):
    """Precios de la lista de ProIsa → precio de la caja (ITBIS incluido).

    Reglas acordadas con el dueño (01/10/2026):
    * `precio_rd` viene SIN ITBIS: ×1.18 da el precio redondo de la caja
      (Michelob 211.86 → 250, Coca 591 46.61 → 55). Se redondea al peso.
    * Si el número ya es redondo (múltiplo de 5) y ×1.18 NO lo es, ya es el
      precio final y va tal cual (Cheetos Grande 145, Princesa 15): 24 casos.
    * Vacío o 1.00 (relleno del sistema viejo) = sin precio.
    * Por código; si el código del producto no está, por nombre EXACTO contra
      las filas cuyo código no es de ningún artículo del inventario (las
      paletas y copas que ProIsa pasó a su código de barras). Ese código
      nuevo, si es GTIN, queda como código de barras del producto.
    """
    path = Path(path).expanduser()
    with open(path, encoding='utf-8-sig', newline='') as fh:
        filas = [{(k or '').strip().lower(): (v or '').strip() for k, v in r.items()}
                 for r in csv.DictReader(fh)]
    if not filas or not {'codigo', 'descripcion', 'precio_rd'} <= set(filas[0]):
        raise SystemExit(f'{path.name}: hacen falta las columnas codigo, descripcion y precio_rd.')
    lista = {}
    for r in filas:
        if r['codigo'] in lista:
            raise SystemExit(f'{path.name}: el código {r["codigo"]} sale dos veces.')
        lista[r['codigo']] = r
    libres = {}
    for code, r in lista.items():
        if code not in conocidos:
            libres.setdefault(_norm(r['descripcion']), []).append(r)

    precios, barcodes, revision = {}, {}, []
    for p in productos:
        r, via = lista.get(p['code']), 'código'
        if r is None:
            cand = libres.get(_norm(p['clean'])) or libres.get(_norm(p['name'])) or []
            r, via = (cand[0], 'nombre') if len(cand) == 1 else (None, '')
            if r is not None:
                bc = barcode_for(r['codigo'])
                if bc:
                    barcodes[p['code']] = bc
        v = float(r['precio_rd']) if r and r['precio_rd'] else None
        if v is None or v <= 1.0:
            final, regla = None, ('no está en la lista' if r is None else
                                  'sin precio en la lista' if v is None else 'precio 1.00 (relleno)')
        elif _redondo(v) and not _redondo(v * 1.18):
            final, regla = round(v, 2), 'ya es precio final (tal cual)'
        else:
            final = float((Decimal(str(v)) * Decimal('1.18')).quantize(Decimal('1'), ROUND_HALF_UP))
            regla = 'sin ITBIS × 1.18, al peso'
        precios[p['code']] = final
        cost = p['cost_out']
        revision.append({
            'categoria': p['group'], 'codigo': p['code'], 'nombre': p['clean'],
            'codigo_proisa': r['codigo'] if r else '', 'emparejado_por': via if r else '',
            'precio_proisa': r['precio_rd'] if r else '', 'precio_caja': '' if final is None else f'{final:.2f}',
            'regla': regla, 'costo': '' if cost is None else f'{cost:.2f}',
            'nota': ('vende por DEBAJO del costo' if final and cost and final / 1.18 < cost else ''),
        })

    ya_finales = sum(1 for x in revision if x['regla'].startswith('ya es'))
    if ya_finales != esperados:
        raise SystemExit(f'{path.name}: {ya_finales} precios parecen ya finales y se esperaban '
                         f'{esperados}. Revisar con el dueño antes de cargar.')
    return precios, barcodes, revision, path.name


UPDATE_TPL = """-- ============================================================================
-- ESTANCIA NUEVA SPORT · CAFETERÍA — PRECIOS (solo UPDATE)
-- Business {bid}
--
-- ⚠ ARCHIVO GENERADO por build_import_cafeteria.py. No lo edites a mano.
-- Precios: {fuente}.
--
-- Los productos ya están cargados (IMPORT_CAFETERIA.sql); esto solo les pone
-- precio. Por producto (sku = código del sistema viejo):
--   · price = precio de la caja, ITBIS incluido
--   · barcode = el código de barras nuevo, solo en los {n_bc} que ProIsa pasó a su
--     EAN (paletas y copas); los demás conservan el suyo
--   · is_active = true si estaba apagado en $0 (esperando precio); uno con
--     precio que apagaron a mano NO se reactiva
-- Se puede correr más de una vez. La OFERTA 3x2 (1219) no se toca.
-- Resultado esperado: con_precio = {n}, activos = {n}.
-- ============================================================================
with precios (codigo, precio, barcode) as (
  values
{values}
),
cambiados as (
  update public.menu_items mi
     set price      = p.precio,
         barcode    = coalesce(p.barcode, mi.barcode),
         is_active  = case when mi.price = 0 and not mi.is_active then true
                           else mi.is_active end,
         updated_at = now()
    from precios p
   where mi.business_id = '{bid}'
     and mi.sku = p.codigo
  returning mi.is_active
)
select count(*)                          as con_precio,
       count(*) filter (where is_active) as activos,
       {n}                               as esperado
from cambiados;
"""


def write_update(productos, barcodes_nuevos, fuente, out):
    con = [r for r in productos if r['vendible'] and r['price'] is not None]
    vals = []
    for i, r in enumerate(con):
        bc = barcodes_nuevos.get(r['code'])
        bc_sql = q(bc) if bc else 'null'
        if i == 0:
            vals.append(f"  ({q(r['code'])}, {r['price']:.2f}::numeric, {bc_sql}::text)")
        else:
            vals.append(f"  ({q(r['code'])}, {r['price']:.2f}, {bc_sql})")
    sql = UPDATE_TPL.format(bid=BID, fuente=fuente, n=len(con),
                            n_bc=len(barcodes_nuevos), values=',\n'.join(vals))
    (out / 'UPDATE_precios_cafeteria.sql').write_text(sql, encoding='utf-8')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--precios')
    ap.add_argument('--proisa')
    ap.add_argument('--out', default=str(HERE))
    args = ap.parse_args()
    if args.precios and args.proisa:
        raise SystemExit('Usa --precios o --proisa, no los dos.')

    productos, insumos, tienda, decidir, *_ = prepare()

    precios, fuente, revision, barcodes = ({}, None, None, {})
    if args.proisa:
        conocidos = {r['code'] for r in productos + insumos + tienda + decidir}
        precios, barcodes, revision, fuente = leer_proisa(args.proisa, productos, conocidos)
        for r in productos:
            if r['code'] in barcodes and not r['barcode']:
                r['barcode'] = barcodes[r['code']]
        fuente += ' (sin ITBIS ×1.18 al peso; 24 ya finales)'
    elif args.precios:
        precios, fuente = leer_precios(args.precios)
        conocidos = {r['code'] for r in productos}
        ajenos = sorted(set(precios) - conocidos)
        if ajenos:
            raise SystemExit('Códigos con precio que no son productos de la cafetería: '
                             + ', '.join(ajenos))

    for cat in CATEGORIAS:
        grupo = [r for r in productos if r['group'] == cat]
        for i, r in enumerate(sorted(grupo, key=lambda r: r['clean']), 1):
            r['pos'] = i
    for r in productos:
        r['price'] = precios.get(r['code'])
        r['vendible'] = r['tipo'] != 'Oferta'
        r['inventariable'] = r['vendible']
        r['area'] = 'cocina' if r['group'] in COCINA else 'bar'
        r['is_bev'] = r['group'] in BEBIDAS
        r['activo'] = r['vendible'] and r['price'] is not None

    def num(x, dec):
        return 'null' if x is None else f'{x:.{dec}f}'

    def b(x):
        return 'true' if x else 'false'

    lotes = []
    for i in range(0, len(productos), 200):
        valores = ',\n'.join(
            '  ({}, {}, {}, {}, {}, {}, {}, {}, {}, {}, {})'.format(
                q(r['code']), q(r['clean']), q(r['group']), num(r['price'], 2),
                num(r['cost_out'], 4), q(r['barcode']) if r['barcode'] else 'null',
                b(r['is_bev']), b(r['vendible']), b(r['inventariable']),
                q(r['area']), r['pos'])
            for r in productos[i:i + 200])
        lotes.append(
            f'-- filas {i + 1}–{min(i + 200, len(productos))}\n'
            'insert into _pen (codigo, name, categoria, price, cost, barcode, '
            'is_bev, vendible, inventariable, area, posicion) values\n'
            + valores + ';')

    con_precio = [r for r in productos if r['activo']]
    subs = {
        '__DATA__': '\n\n'.join(lotes),
        '__INSUMOS__': ',\n'.join(
            '  ({}, {}, {}, {})'.format(q(r['code']), q(r['clean']),
                                        num(r['cost_out'], 4),
                                        q(r['barcode']) if r['barcode'] else 'null')
            for r in insumos),
        '__CATEGORIAS__': ',\n'.join(
            f'  ({q(c)}, {(i + 1) * 10})' for i, c in enumerate(CATEGORIAS)),
        '__CAT_NAMES_LOWER__': ', '.join(q(c.lower()) for c in CATEGORIAS),
        '__CODES__': ','.join(r['code'] for r in productos),
        '__INS_CODES__': ','.join(r['code'] for r in insumos),
        '__BID__': BID,
        '__FUENTE_PRECIOS__': (f'{fuente} ({len(con_precio)} con precio)' if fuente
                               else 'TODAVÍA NO — todo entra inactivo en RD$0'),
        '__N_TOTAL__': str(len(productos)),
        '__N_CON_PRECIO__': str(len(con_precio)),
        '__N_INVENTARIABLES__': str(sum(r['inventariable'] for r in productos)),
        '__N_INSUMOS__': str(len(insumos)),
        '__N_BARCODE__': str(sum(1 for r in productos if r['barcode'])),
        '__N_CATEGORIAS__': str(len(CATEGORIAS)),
        '__N_COCINA__': str(sum(r['area'] == 'cocina' for r in productos)),
        '__N_BAR__': str(sum(r['area'] == 'bar' for r in productos)),
    }

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    for tpl, name in (('_tpl_import_cafeteria.sql', 'IMPORT_CAFETERIA.sql'),
                      ('_tpl_rollback_cafeteria.sql', '99_rollback_cafeteria.sql')):
        sql = (HERE / tpl).read_text(encoding='utf-8')
        for k, v in subs.items():
            sql = sql.replace(k, v)
        sobrantes = re.findall(r'__[A-Z_]+__', sql.replace('__IN_TRANSIT__', ''))
        if sobrantes:
            sys.exit(f'{tpl}: placeholders sin reemplazar: {sorted(set(sobrantes))}')
        (out / name).write_text(sql, encoding='utf-8')

    if fuente:
        write_update(productos, barcodes, fuente, out)

    if revision is not None:
        with open(out / 'precios_cafeteria_revision.csv', 'w', newline='', encoding='utf-8') as fh:
            w = csv.DictWriter(fh, fieldnames=list(revision[0]))
            w.writeheader()
            w.writerows(revision)

    print(f"{subs['__N_TOTAL__']} productos ({subs['__N_CON_PRECIO__']} con precio → activos) · "
          f"{subs['__N_INVENTARIABLES__']} inventariables · {subs['__N_INSUMOS__']} insumos de cocina · "
          f"COCINA {subs['__N_COCINA__']} / BAR {subs['__N_BAR__']} · "
          f"{subs['__N_BARCODE__']} con código de barras · {subs['__N_CATEGORIAS__']} categorías")


if __name__ == '__main__':
    main()
